import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank)

img = "screenshots/01_patients_list.png"
if os.path.exists(img):
    slide.shapes.add_picture(img, Inches(1.0), Inches(1.0), width=Inches(8.0))
    prs.save("screenshots/test_embed.pptx")
    print("SUCCESS: Embedded picture into test_embed.pptx")
else:
    print("Image not found:", img)
