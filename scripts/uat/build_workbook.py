"""Build reports/uat/UAT_RESULTS_ASTER.xlsx: the QA workbook with the retest status and evidence filled in.

    python scripts/uat/build_workbook.py [overrides.json]

Reads reports/uat/uat_results.json (written by run_uat.py) and the original E2E_WORKFLOW_TEST_SCENARIOS.xlsx. The
optional overrides file ({id: {"status": "Pass", "note": "..."}}) records cases proven some other way than the API run
(for example a screen check, or a database access-rule listing).
"""
import json
import os
import sys

import openpyxl
from openpyxl.worksheet.table import Table

ROOT = os.path.join(os.path.dirname(__file__), "..", "..")
SRC = os.path.join(ROOT, "reports", "uat", "E2E_WORKFLOW_TEST_SCENARIOS.xlsx")
RES = os.path.join(ROOT, "reports", "uat", "uat_results.json")
OUT = os.path.join(ROOT, "reports", "uat", "UAT_RESULTS_ASTER.xlsx")

results = {r["id"]: r for r in json.load(open(RES, encoding="utf-8"))}
overrides = json.load(open(sys.argv[1], encoding="utf-8")) if len(sys.argv) > 1 else {}

wb = openpyxl.load_workbook(SRC)
ws = wb["Test Cases"]
ws.cell(row=1, column=11, value="Retest evidence")
ws.cell(row=1, column=12, value="Retest basis")
for row in range(2, ws.max_row + 1):
    tid = ws.cell(row=row, column=1).value
    if not tid:
        continue
    o = overrides.get(tid)
    r = results.get(tid)
    if o:
        status, evidence, basis = o["status"], o["note"], o.get("basis", "Verified on the live Aster tenant")
    elif r:
        status = "Pass" if r["status"] == "PASS" else "Fail"
        evidence, basis = r["observed"], "Automated API test on the live Aster tenant"
    else:
        status, evidence, basis = "Blocked", "No result recorded.", ""
    ws.cell(row=row, column=9, value=status)
    ws.cell(row=row, column=11, value=evidence[:900])
    ws.cell(row=row, column=12, value=basis)

for t in ws.tables.values():
    t.ref = f"A1:L{ws.max_row}"
    break
wb.save(OUT)
counts = {}
for row in range(2, ws.max_row + 1):
    s = ws.cell(row=row, column=9).value
    counts[s] = counts.get(s, 0) + 1
print(OUT, counts)
